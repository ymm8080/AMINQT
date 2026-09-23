# -*- coding: utf-8 -*-
"""首板点名页 + 板前哨页 数据层测试 (scripts/_firstboard_pages.py, 0922)。

覆盖: 事件口径(首板=板且前10日无板) / wr1=t-1 / 冠军格 / promo 修正版无前视 /
深睡签名各腿 / 点火旗 T1 / 板前哨滚动20日命中日志(多行+显示过滤+排序+WORM后台表+
停牌缺今日行) / 模型训练-落盘-服务 roundtrip / GENIOUS 多页写出+页底脚注。
"""

from __future__ import annotations

import importlib
import logging
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

fbp = importlib.import_module("scripts._firstboard_pages")

DATES = pd.bdate_range("2026-01-01", periods=90)


@pytest.fixture(autouse=True)
def _no_lhb_seats(monkeypatch):
    """0923: LHB 席位数据默认空表 (勿读生产 parquet, 保持测试密闭); 标注测试内自行注入缓存."""
    monkeypatch.setattr(fbp, "_LHB_DAILY_CACHE", pd.DataFrame())


def mk_panel(rows: dict[str, dict[str, list]]) -> pd.DataFrame:
    """rows: symbol -> 覆盖列(k=v 列表, 长度 90)。其余列给恒定默认值。"""
    n = len(DATES)
    frames = []
    defaults = dict(
        open=10.0,
        high=10.5,
        low=9.5,
        close=10.0,
        pre_close=10.0,
        pctChg=0.0,
        turnover_rate=1.0,
        volume_ratio=1.0,
        amount=1e8,
        circ_mv=30.0,
        free_float_turnover_rate=5.0,
        winner_ratio=0.30,
        sw_l2_name="X1",
        lhb_net_buy=0.0,
        lhb_inst_buy=0.0,
    )
    for sym, ov in rows.items():
        d = {"symbol": sym, "date": DATES}
        for col, default in defaults.items():
            v = ov.get(col, default)
            d[col] = v if isinstance(v, list) else [v] * n
        frames.append(pd.DataFrame(d))
    df = pd.concat(frames, ignore_index=True)
    df["date"] = pd.to_datetime(df["date"])
    return df.sort_values(["symbol", "date"]).reset_index(drop=True)


def _set(df, sym, iloc, **kv):
    for k, v in kv.items():
        df.loc[(df["symbol"] == sym) & (df["date"] == DATES[iloc]), k] = v


class TestEventFeatures:
    def test_firstboard_event_and_prior10_exclusion(self):
        ov = {"pctChg": [0.0] * len(DATES)}
        df = mk_panel({"600001": ov, "600002": ov})
        # 600001: 第70日首板 → 事件
        _set(df, "600001", 70, pctChg=10.0)
        # 600002: 第60日板 + 第70日再板 → 只有第60日是事件, 第70日被 prior10 排除
        _set(df, "600002", 60, pctChg=10.0)
        _set(df, "600002", 70, pctChg=10.0)
        ev = fbp.build_event_features(df)
        a = ev[ev["symbol"] == "600001"]
        b = ev[ev["symbol"] == "600002"]
        assert len(a) == 1 and a.iloc[0]["date"] == DATES[70]
        assert len(b) == 1 and b.iloc[0]["date"] == DATES[60]

    def test_wr1_is_t_minus_1_winner_ratio(self):
        ov = {
            "pctChg": [0.0] * len(DATES),
            "winner_ratio": [0.30 + i * 0.001 for i in range(len(DATES))],
        }
        df = mk_panel({"600003": ov})
        _set(df, "600003", 70, pctChg=10.0)
        ev = fbp.build_event_features(df)
        assert ev.iloc[0]["wr1"] == pytest.approx(0.30 + 69 * 0.001)

    def test_champion_flag_yizi_and_wr(self):
        ov = {"pctChg": [0.0] * len(DATES), "winner_ratio": [0.70] * len(DATES)}
        df = mk_panel({"600004": ov})
        # 一字: low 贴涨停价 (pre_close 10 → limit 11.0); 且板日
        _set(df, "600004", 70, pctChg=10.0, low=11.0, open=11.0, high=11.0, close=11.0)
        ev = fbp.build_event_features(df)
        row = ev.iloc[0]
        assert bool(row["yizi"]) and bool(row["champion"])
        # 对照: 低获利盘 → 非冠军
        df2 = mk_panel(
            {
                "600005": {
                    "pctChg": [0.0] * len(DATES),
                    "winner_ratio": [0.40] * len(DATES),
                }
            }
        )
        _set(df2, "600005", 70, pctChg=10.0, low=11.0, open=11.0, high=11.0, close=11.0)
        ev2 = fbp.build_event_features(df2)
        assert bool(ev2.iloc[0]["yizi"]) and not bool(ev2.iloc[0]["champion"])

    def test_label_nan_at_tail(self):
        ov = {"pctChg": [0.0] * len(DATES)}
        df = mk_panel({"600006": ov})
        _set(
            df, "600006", len(DATES) - 2, pctChg=10.0
        )  # 剩1行: k2 需5日→NaN, next 需1日→0.0
        ev = fbp.build_event_features(df)
        assert ev.iloc[0]["k2"] != ev.iloc[0]["k2"]  # NaN
        assert ev.iloc[0]["k3"] != ev.iloc[0]["k3"]  # k3 需3日 → 同样 NaN
        assert ev.iloc[0]["next_board"] == 0.0

    def test_k3_label_window3(self):
        # 0923: k3 = D0+1..D0+3 内任一板 (同 k2 式窗口 3) — 板在 D0+4 只算 k2 不算 k3
        df = mk_panel(
            {
                "600007": {"pctChg": [0.0] * len(DATES)},
                "600008": {"pctChg": [0.0] * len(DATES)},
            }
        )
        _set(df, "600007", 70, pctChg=10.0)
        _set(df, "600007", 73, pctChg=10.0)  # D0+3 再板 → k3=1
        _set(df, "600008", 70, pctChg=10.0)
        _set(df, "600008", 74, pctChg=10.0)  # D0+4 再板 → k3=0, k2=1
        ev = fbp.build_event_features(df)
        a = ev[ev["symbol"] == "600007"].iloc[0]
        b = ev[ev["symbol"] == "600008"].iloc[0]
        assert a["k3"] == 1.0 and a["k2"] == 1.0
        assert b["k3"] == 0.0 and b["k2"] == 1.0


class TestPromoNoLookahead:
    def _tiny(self, y_boards_d2: float):
        dates = pd.bdate_range("2026-03-02", periods=3)
        rows = []
        for sym, boards in (
            ("600101", (1.0, 1.0, 0.0)),
            ("600102", (1.0, 0.0, y_boards_d2)),
        ):
            for i, dt in enumerate(dates):
                rows.append(
                    dict(
                        symbol=sym,
                        date=dt,
                        pctChg=10.0 if boards[i] else 0.0,
                        pre_close=10.0,
                        high=11.0,
                        close=10.5,
                    )
                )
        df = pd.DataFrame(rows)
        df["date"] = pd.to_datetime(df["date"])
        return df.sort_values(["symbol", "date"]).reset_index(drop=True)

    def test_first_day_zero_and_ratio(self):
        eco = fbp._ecology(self._tiny(0.0))
        p = eco.set_index("date")["promo_yd"]
        assert p.iloc[0] == 0.0  # 首日无 D-1
        # d1: 昨日板 {101,102}, 今日板 {101} → 晋级率 1/2
        assert p.iloc[1] == pytest.approx(0.5)

    def test_d2_changes_do_not_leak_into_d1(self):
        # W22 核心: D 日的 promo 不得随 D+1 内容变化 (原漏版违反此性质)
        p0 = fbp._ecology(self._tiny(0.0)).set_index("date")["promo_yd"]
        p1 = fbp._ecology(self._tiny(1.0)).set_index("date")["promo_yd"]
        assert p0.iloc[1] == p1.iloc[1]


class TestDeepSleepSig:
    def _sleep_sym(self, wr_delta=0.15, t5=0.5, t20_head=2.0):
        n = len(DATES)
        wr = [0.30] * 70 + [0.30 + wr_delta * (i - 69) / (n - 70) for i in range(70, n)]
        turn = [1.0] * 70 + [t20_head] * 15 + [t5] * 5
        return {"winner_ratio": wr, "turnover_rate": turn}

    def _with_pulse(self, ov):
        ov = dict(ov)
        ov["pctChg"] = [0.0] * len(DATES)
        ov["volume_ratio"] = [1.0] * len(DATES)
        # t=89, k=30 → 第59日脉冲 (涨≥4% ∩ 量比≥1.8); 其后无板
        ov["pctChg"][59] = 5.0
        ov["volume_ratio"][59] = 2.0
        return ov

    def test_sig_true_on_pattern(self):
        df = mk_panel({"600201": self._with_pulse(self._sleep_sym())})
        s = fbp.build_sig(df)
        assert bool(s.iloc[-1]["sig"])

    def test_sig_requires_chip_leg(self):
        df = mk_panel({"600202": self._with_pulse(self._sleep_sym(wr_delta=0.05))})
        s = fbp.build_sig(df)
        assert not bool(s.iloc[-1]["sig"])

    def test_sig_requires_squeeze_leg(self):
        df = mk_panel({"600203": self._with_pulse(self._sleep_sym(t5=1.5))})
        s = fbp.build_sig(df)
        assert not bool(s.iloc[-1]["sig"])

    def test_board_day_excluded(self):
        ov = self._with_pulse(self._sleep_sym())
        ov["pctChg"][-1] = 10.0  # 当日板 → sig 必须关
        df = mk_panel({"600204": ov})
        s = fbp.build_sig(df)
        assert not bool(s.iloc[-1]["sig"])

    def test_no_pulse_no_sig(self):
        ov = self._sleep_sym()
        ov["pctChg"] = [0.0] * len(DATES)
        ov["volume_ratio"] = [1.0] * len(DATES)
        df = mk_panel({"600205": ov})
        s = fbp.build_sig(df)
        assert not bool(s.iloc[-1]["sig"])


class TestFireFlags:
    def test_t1_first_ge2_in_5d(self):
        pct = [0.0] * len(DATES)
        pct[-1] = 3.0
        df = mk_panel({"600301": {"pctChg": pct}})
        s = fbp.build_sig(df)
        assert bool(s.iloc[-1]["fire_t1"]) and bool(s.iloc[-1]["fire_t2"])

    def test_t1_off_if_recent_ge2(self):
        pct = [0.0] * len(DATES)
        pct[-4] = 2.5
        pct[-1] = 3.0
        df = mk_panel({"600302": {"pctChg": pct}})
        s = fbp.build_sig(df)
        assert not bool(s.iloc[-1]["fire_t1"])


class TestPreboardRolling:
    """serve_preboard 滚动20日命中日志语义 (v4 0923: 显示=仅今日点火旗股, 次数退出筛选)。

    构造法: wr 台阶 0.30→0.42 (CHIP 腿 wr20=0.12≥0.10 在台阶日起 20 日内成立);
    turn 前段 2.0 后段 0.4 (SQUEEZE 腿 t5=0.4≤0.75×t20 在 85 日后成立);
    第50日脉冲 (PULSE 腿, 85..89 日对应 k=35..39 ∈[10,40]); close 恒定 (hold/横盘腿)。
    wr_step_at 控制每股 sig 日集合: 85 → sig{85..89}(次数5) / 88 → {88,89}(次数2) /
    89 → {89}(次数1)。今日 pctChg[-1]=3.0 → 点火 T1+T2 (sig 不受影响, 3.0<9.5 非板)。
    """

    def _pb_sym(self, wr_step_at=85):
        n = len(DATES)
        ov = {
            "winner_ratio": [0.30] * wr_step_at + [0.42] * (n - wr_step_at),
            "turnover_rate": [2.0] * 75 + [0.4] * (n - 75),
            "pctChg": [0.0] * n,
            "volume_ratio": [1.0] * n,
        }
        ov["pctChg"][50] = 5.0
        ov["volume_ratio"][50] = 2.0
        return ov

    def test_multirow_and_fire_only_display(self):
        ova = self._pb_sym(85)
        ova["pctChg"][-1] = 3.0  # 今日点火 → v4 唯一显示条件
        df = mk_panel(
            {"600501": ova, "600502": self._pb_sym(89), "600504": self._pb_sym(88)}
        )
        page = fbp.serve_preboard(df)
        # 股A 今日点火 → 全部命中日各一行; 两次数列=近10/20交易日命中总数 (逐行同值)
        a = page[page["代码"] == "600501"]
        assert len(a) == 5
        assert (a["次数10日"] == 5).all() and (a["次数20日"] == 5).all()
        assert set(a["击中日期"]) == {
            DATES[i].strftime("%Y-%m-%d") for i in range(85, 90)
        }
        # v4: 无点火 → 整组不显示 (次数1 股B / 次数≥2 股D 都进不了 Excel)
        assert set(page["代码"]) == {"600501"}

    def test_fire_flag_rescues_count1(self):
        ov = self._pb_sym(89)
        ov["pctChg"][-1] = 3.0  # T1: 今日≥2 且前5日无≥2; T2: 涨2~7 → T1+T2
        ov["winner_ratio"][-1] = 0.70  # 通道优先腿1: 今日获利盘≥0.65
        ov["lhb_net_buy"] = [0.0] * 85 + [
            1.0
        ] * 5  # 通道优先腿2: 20日内LHB净买+机构双旗
        ov["lhb_inst_buy"] = [0.0] * 85 + [1.0] * 5
        page = fbp.serve_preboard(mk_panel({"600503": ov}))
        assert len(page) == 1
        r = page.iloc[0]
        assert (
            r["代码"] == "600503" and r["点火旗"] == "T1+T2" and r["通道优先"] == "优先"
        )
        assert r["次数10日"] == 1 and r["次数20日"] == 1

    def test_sort_fire_then_date_then_count(self):
        ova = self._pb_sym(85)
        ova["pctChg"][-1] = 3.0  # T1+T2, 次数5
        ovc = self._pb_sym(89)
        ovc["pctChg"][-1] = 3.0  # T1+T2, 次数1
        ovd = self._pb_sym(88)
        ovd["pctChg"][-2] = 2.5  # 前5日已有≥2 → T1 关
        ovd["pctChg"][-1] = 3.0  # 涨2~7 → 仅 T2 (v4 无点火不显示, 须给火)
        df = mk_panel({"600501": ova, "600503": ovc, "600504": ovd})
        page = fbp.serve_preboard(df)
        d = lambda i: DATES[i].strftime("%Y-%m-%d")  # noqa: E731
        got = list(zip(page["代码"], page["击中日期"]))
        # 点火等级高在前 (T1+T2 > T2); 同等级内 击中日期新在前; 同日期 次数10日大在前 (A=5 > C=1)
        assert got == [
            ("600501", d(89)),
            ("600503", d(89)),
            ("600501", d(88)),
            ("600501", d(87)),
            ("600501", d(86)),
            ("600501", d(85)),
            ("600504", d(89)),
            ("600504", d(88)),
        ]

    def test_channel_priority_sorts_first(self):
        """通道优先行必须置顶 (页首列; PB_LEGEND『通道优先=…(置顶)』『它是排序提示』)。

        构造: 优先股B 点火仅 T2 (等级最低) 且命中集合与A同 ⇒ 旧键 (_fr 优先) 下 B 五行
        必排到 A 五行之后; 修复后 (_pass 优先) B 五行置顶。本测试在修复前必失败。
        """
        ova = self._pb_sym(85)
        ova["pctChg"][-1] = 3.0  # T1+T2 (最高等级), 非优先
        ovb = self._pb_sym(85)
        ovb["pctChg"][-2] = 2.5  # 前5日已≥2 → T1 关
        ovb["pctChg"][-1] = 3.0  # 涨2~7 → 仅 T2 (最低等级)
        ovb["winner_ratio"][-1] = 0.70  # 通道优先腿1: 今日获利盘≥0.65
        ovb["lhb_net_buy"] = [0.0] * 85 + [1.0] * 5  # 腿2: 20日内LHB净买+机构双旗
        ovb["lhb_inst_buy"] = [0.0] * 85 + [1.0] * 5
        page = fbp.serve_preboard(mk_panel({"600501": ova, "600503": ovb}))
        flags = page["通道优先"].tolist()
        assert flags.count("优先") == 5
        # 全部优先行连续置顶, 非优先行全部在后
        assert flags == ["优先"] * flags.count("优先") + [""] * flags.count("")
        assert set(page.loc[page["通道优先"] == "优先", "代码"]) == {"600503"}
        assert page["点火旗"].iloc[0] == "T2"  # 置顶的恰是最低点火等级(优先)股

    def test_record_csv_full_and_worm(self, tmp_path):
        df = mk_panel({"600501": self._pb_sym(85), "600502": self._pb_sym(89)})
        p = tmp_path / "preboard_watch_hits_test.csv"
        fbp.serve_preboard(df, record_csv=p)
        assert p.exists()
        rec = pd.read_csv(
            p, encoding="utf-8-sig", keep_default_na=False, dtype={"代码": str}
        )
        assert rec.columns.tolist() == [
            "代码",
            "首次击中",
            "最近击中",
            "次数10日",
            "次数20日",
            "今日点火旗",
            "今日涨幅",
            "今日获利盘",
            "通道优先",
        ]
        # 后台表全量: 含被显示过滤掉的股B
        assert sorted(rec["代码"]) == ["600501", "600502"]
        ra = rec[rec["代码"] == "600501"].iloc[0]
        assert (ra["首次击中"], ra["最近击中"]) == (
            DATES[85].strftime("%Y-%m-%d"),
            DATES[89].strftime("%Y-%m-%d"),
        )
        assert ra["次数10日"] == 5 and ra["次数20日"] == 5
        rb = rec[rec["代码"] == "600502"].iloc[0]
        assert rb["次数10日"] == 1 and rb["今日点火旗"] == ""
        assert float(rb["今日涨幅"]) == pytest.approx(0.0)  # pctChg=0 → 小数 0.0
        # WORM: 同名重跑不覆盖, 生成 __HHMMSS 变体, 原文件字节不变
        raw1 = p.read_bytes()
        time.sleep(1.1)  # 避开同秒 HHMMSS 碰撞
        fbp.serve_preboard(df, record_csv=p)
        assert p.read_bytes() == raw1
        assert len(list(tmp_path.glob("preboard_watch_hits_test__*.csv"))) == 1

    def test_suspended_missing_today_row(self):
        ova = self._pb_sym(85)
        ova["pctChg"][-1] = 3.0  # 股A 今日点火 (保留今日行)
        df = mk_panel({"600501": ova, "600502": self._pb_sym(85)})
        df = df[~((df["symbol"] == "600502") & (df["date"] == DATES[-1]))].reset_index(
            drop=True
        )
        page = fbp.serve_preboard(df)
        # 股A 点火 → 5 个命中日各一行
        a = page[page["代码"] == "600501"]
        assert len(a) == 5 and (a["点火旗"] == "T1+T2").all()
        # 股B 缺今日行(停牌) → 点火旗必空 → v4 整组不显示 (次数10日=4 也救不回)
        assert set(page["代码"]) == {"600501"}


@pytest.fixture(scope="module")
def trained(tmp_path_factory):
    rng = np.random.default_rng(42)
    n_sym, n_day = 16, 650
    dates = pd.bdate_range("2024-01-01", periods=n_day)
    rows = []
    for k in range(n_sym):
        wr = 0.3 + 0.4 * rng.random(n_day).cumsum() / n_day
        pct = rng.normal(0, 2.0, n_day)
        is_board = rng.random(n_day) < 0.04  # 肥尾注入板日, 否则正态无板
        pct[is_board] = rng.uniform(9.6, 11.0, int(is_board.sum()))
        if k == 0:  # 末日保证一个干净首板事件 (serve 出页必须有当日事件)
            pct[-12:-1] = 0.0
            pct[-1] = 10.0
        for i, dt in enumerate(dates):
            rows.append(
                dict(
                    symbol=f"6004{k:02d}",
                    date=dt,
                    open=10.0,
                    high=10.5,
                    low=9.5,
                    close=10.0,
                    pre_close=10.0,
                    pctChg=float(pct[i]),
                    turnover_rate=1 + rng.random(),
                    volume_ratio=1 + rng.random(),
                    amount=1e8,
                    circ_mv=30.0,
                    free_float_turnover_rate=5.0,
                    winner_ratio=float(np.clip(wr[i], 0, 1)),
                    sw_l2_name="X1",
                    lhb_net_buy=0.0,
                    lhb_inst_buy=0.0,
                )
            )
    df = pd.DataFrame(rows).sort_values(["symbol", "date"]).reset_index(drop=True)
    df["date"] = pd.to_datetime(df["date"])
    pq_path = tmp_path_factory.mktemp("fbp") / "synthetic_panel.parquet"
    df.to_parquet(pq_path, index=False)
    out = tmp_path_factory.mktemp("models") / "firstboard"
    res = fbp.train_and_save(out_dir=out, panel_path=str(pq_path), expect=None)
    return out, df, res


class TestModelRoundtrip:
    def test_files_written(self, trained):
        out, _, _ = trained
        assert (out / "booster_k2.txt").exists()
        assert (out / "booster_k3.txt").exists()
        assert (out / "booster_next.txt").exists()
        import json

        m = json.loads((out / "meta.json").read_text(encoding="utf-8"))
        assert m["features"] == fbp.FEATS and len(m["sw_l2_categories"]) >= 1
        # 0923: k3 头验收线登记进 meta (真实验收断言在 --train 于真面板上执行)
        assert "te_top1_k3" in m and "te_auc_k3" in m

    def test_k3_reload_reproduction_registered(self, trained):
        # 训练即 reload 复现: k3 指标算出并原样登记进 meta (防落盘口径分裂)
        out, _, res = trained
        import json

        m = json.loads((out / "meta.json").read_text(encoding="utf-8"))
        assert m["te_top1_k3"] == pytest.approx(round(res["k3@1"], 1), abs=1e-9)
        if res["auc_k3"] == res["auc_k3"]:  # NaN (合成TE单类) 跳过
            assert m["te_auc_k3"] == pytest.approx(round(res["auc_k3"], 4), abs=1e-9)
        else:
            assert m["te_auc_k3"] != m["te_auc_k3"]

    def test_expect_te_k3_lines_pinned(self):
        # 0923 用户令断言线: p_k3 Top1=39.8±0.5, AUC_k3=0.5716±0.005 (rank_calib 实测)
        assert fbp._EXPECT_TE["k3@1"] == 39.8
        assert fbp._EXPECT_TE["auc_k3"] == 0.5716
        assert fbp._EXPECT_TE["tol_k3_hits"] == 0.5
        assert fbp._EXPECT_TE["tol_k3_auc"] == 0.005
        # 0923 刷新 k2 两线: 旧线 46.1/42.0 建于 0922 V3 面板重建(BJ剔除)前 = 旧基线;
        # 当前帧生产 booster 实读 45.78/41.77 → 45.8/41.8 (auc_next 线不动)
        assert (
            fbp._EXPECT_TE["k2@1"],
            fbp._EXPECT_TE["k2@3"],
            fbp._EXPECT_TE["auc_next"],
        ) == (45.8, 41.8, 0.5911)

    def test_serve_roundtrip_probs_in_range(self, trained, monkeypatch):
        out, df, _ = trained
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})  # 测试不触网
        page = fbp.serve_firstboard(df, models_dir=out)
        assert {
            "排名",
            "冠军格",
            "代码",
            "名称",
            "T+3板概率",
            "T+5板概率",
            "校准档位",
            "一字",
            "次日一字风险",
            "板块涨停数",
            "市场涨停数",
        } <= set(page.columns)
        got = page["T+3板概率"].dropna()
        assert len(got) and got.between(0, 1).all()
        got5 = page["T+5板概率"].dropna()
        assert len(got5) and got5.between(0, 1).all()
        assert page["排名"].dropna().tolist() == sorted(page["排名"].dropna().tolist())

    def test_preboard_page_columns(self, trained):
        _, df, _ = trained
        page = fbp.serve_preboard(df)
        assert page.columns.tolist() == [
            "通道优先",
            "点火旗",
            "击中日期",
            "次数10日",
            "次数20日",
            "代码",
            "当日涨幅",
            "获利盘",
            "20日获利盘Δ",
            "5日均换手",
            "5日/20日换手比",
        ]
        if len(page):  # 击中日期为文本日期, 两次数列为整数
            assert page["击中日期"].map(lambda v: isinstance(v, str)).all()
            assert pd.api.types.is_integer_dtype(page["次数10日"])
            assert pd.api.types.is_integer_dtype(page["次数20日"])


class TestCalibTier:
    def test_bucket_boundaries(self):
        p = pd.Series([0.50, 0.40, 0.399, 0.30, 0.299, 0.20, 0.199, np.nan])
        t = fbp._calib_tier(p)
        assert t.tolist() == [
            "≥0.4→实测~54%",
            "≥0.4→实测~54%",
            "0.3-0.4→实测~28%",
            "0.3-0.4→实测~28%",
            "0.2-0.3→实测~21%",
            "0.2-0.3→实测~21%",
            "<0.2→实测~18%",
            "",
        ]


def _today_events_df(n=6, champion_sym=None):
    """n 只股票同日(最后一个交易日)首板 — 全部事件可打分 (wr/ret60/eco 齐全)."""
    rows = {}
    for k in range(n):
        sym = f"6006{k:02d}"
        ov = {
            "pctChg": [0.0] * len(DATES),
            "winner_ratio": [0.30 + 0.05 * k] * len(DATES),
        }
        ov["pctChg"][-1] = 10.0
        if sym == champion_sym:  # 一字 + 高获利盘 → 冠军格
            ov["winner_ratio"] = [0.70] * len(DATES)
            for col, v0, v1 in (
                ("open", 10.0, 11.0),
                ("high", 10.5, 11.0),
                ("low", 9.5, 11.0),
                ("close", 10.0, 11.0),
            ):
                ov[col] = [v0] * (len(DATES) - 1) + [v1]
        rows[sym] = ov
    return mk_panel(rows)


class TestServeFirstboardV2:
    """点名页 v2 (0923): 全量按 p_k3 降序 + 校准档位 + 冠军格并集保留 + 一字提示语义."""

    def test_columns_and_pk3_descending(self, trained, monkeypatch):
        out, _, _ = trained
        monkeypatch.setattr(fbp, "_NAME_CACHE", {"600600": "甲股"})
        page = fbp.serve_firstboard(_today_events_df(6), models_dir=out)
        assert (
            page.columns.tolist()
            == [
                "排名",
                "代码",
                "名称",
                "T+3板概率",
                "T+5板概率",
                "校准档位",
                "冠军格",
                "一字",
                "次日一字风险",
                "板前获利盘",
                "板块涨停数",
                "昨日晋级率",
                "市场涨停数",
            ]
            + fbp.LHB_ANNOT_COLS
        )
        assert len(page) == 6  # 全量清单, 不 Top-N 截断
        p3 = page["T+3板概率"]
        assert p3.notna().all()
        assert (p3.diff().dropna() <= 1e-12).all()  # 主排序键降序
        assert page["排名"].tolist() == list(range(1, 7))
        assert page.loc[page["代码"] == "600600", "名称"].iloc[0] == "甲股"

    def test_calib_column_from_tiers(self, trained, monkeypatch):
        out, _, _ = trained
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})
        page = fbp.serve_firstboard(_today_events_df(6), models_dir=out)
        allowed = {
            "≥0.4→实测~54%",
            "0.3-0.4→实测~28%",
            "0.2-0.3→实测~21%",
            "<0.2→实测~18%",
        }
        assert set(page["校准档位"]) <= allowed
        for tier, p in zip(page["校准档位"], page["T+3板概率"]):
            assert tier == fbp._calib_tier(pd.Series([p])).iloc[0]

    def test_champion_kept_when_truncated(self, trained, monkeypatch):
        out, _, _ = trained
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})
        page = fbp.serve_firstboard(
            _today_events_df(8, champion_sym="600607"), models_dir=out, max_rows=5
        )
        # 截 5 行, 冠军格★行豁免保留 (W19 并集规则)
        assert 5 <= len(page) <= 6
        assert (page["冠军格"] == "★").any()
        assert set(page["冠军格"]) <= {"", "★"}

    def test_yizi_is_hint_not_veto_and_risk_consistency(self, trained, monkeypatch):
        out, _, _ = trained
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})
        page = fbp.serve_firstboard(
            _today_events_df(6, champion_sym="600600"), models_dir=out
        )
        yz = page[page["一字"] == "一字"]
        assert len(yz) == 1  # 一字行仍在清单 = 提示非否决
        # 次日一字风险 = 一字 ∩ T+3板概率≥阈值 (仅提示, 页内自洽)
        exp = np.where(
            (page["一字"] == "一字") & (page["T+3板概率"] >= fbp.K3_RISK_TH), "风险", ""
        )
        assert (page["次日一字风险"].to_numpy() == exp).all()
        assert set(page["次日一字风险"]) <= {"", "风险"}


class TestServeFirstboardMissingModels:
    """模型文件缺失 = 静默失败 (p_k3 全 NaN → 排名/校准档位空, 页失效) ⇒ 必须发声。"""

    def test_missing_models_logs_error(self, tmp_path, caplog, monkeypatch):
        # 本测试要求真 ERROR 记录, 不是『只是不抛异常』
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})
        caplog.set_level(logging.ERROR, logger=fbp.log.name)
        page = fbp.serve_firstboard(_today_events_df(2), models_dir=tmp_path)
        errs = [r for r in caplog.records if r.levelno >= logging.ERROR]
        assert errs, "首板模型缺失必须发 ERROR 日志 (否则页面看似正常却零信息)"
        msg = " ".join(r.getMessage() for r in errs)
        assert "排名" in msg and "校准档位" in msg and str(tmp_path) in msg
        assert page["T+3板概率"].isna().all()  # 确认确实走了缺模型分支


class TestLhbSeatAnnotations:
    """0923 LHB 席位标注列: 窗=D-20..D0 含 D0 (面板日历), 元→亿, 无数据留空绝不删行."""

    @staticmethod
    def _daily(rows):
        return pd.DataFrame(
            rows,
            columns=["symbol", "date", "n_rows", "net_sum", "hotseat_n", "inst_any"],
        )

    def test_agg_seat_daily_categories(self):
        st = pd.DataFrame(
            {
                "ts_code": ["000610.SZ", "000610.SZ", "600001.SH", "300001.SZ"],
                "trade_date": ["20260105"] * 4,
                "net_buy": [100_000_000.0, -50_000_000.0, 30_000_000.0, 10.0],
                "category": ["高频游资", "机构专用", "拉萨散户团", "高频游资"],
            }
        )
        d = fbp._agg_seat_daily(st)
        a = d[d["symbol"] == "000610"].iloc[0]
        assert a["net_sum"] == pytest.approx(0.5)  # (1-0.5)亿
        assert a["hotseat_n"] == 1.0 and a["inst_any"] == 1.0 and a["n_rows"] == 2.0
        assert set(d["symbol"]) == {"000610", "600001"}  # 创业板 300 被主板过滤
        b = d[d["symbol"] == "600001"].iloc[0]
        assert b["hotseat_n"] == 0.0 and b["inst_any"] == 0.0

    def test_page_annotations_and_blank_rows(self, trained, monkeypatch):
        out, _, _ = trained
        monkeypatch.setattr(
            fbp,
            "_LHB_DAILY_CACHE",
            self._daily(
                [
                    ("600600", DATES[-1], 3.0, 2.5, 4.0, 1.0),  # D0 当日
                    ("600600", DATES[-3], 2.0, -1.0, 1.0, 0.0),  # 窗内前史
                    ("600601", DATES[-1], 1.0, 0.8, 0.0, 0.0),
                ]
            ),
        )
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})
        page = fbp.serve_firstboard(_today_events_df(3), models_dir=out)
        r0 = page[page["代码"] == "600600"].iloc[0]
        assert r0["LHB上榜日数(20日)"] == 2.0
        assert r0["LHB净买亿(20日)"] == pytest.approx(1.5)
        assert r0["游资席位数(20日)"] == 5.0
        assert r0["机构在场(20日)"] == "在"
        assert r0["D0游资席位"] == 4.0 and r0["D0净买亿"] == pytest.approx(2.5)
        r1 = page[page["代码"] == "600601"].iloc[0]
        assert r1["机构在场(20日)"] == "" and r1["D0游资席位"] == 0.0
        # 无 LHB 数据股: 行仍在, 标注列空 (观察清单绝不删行)
        r2 = page[page["代码"] == "600602"].iloc[0]
        assert np.isnan(r2["LHB上榜日数(20日)"]) and r2["机构在场(20日)"] == ""
        assert len(page) == 3

    def test_window_is_d20_to_d0(self, trained, monkeypatch):
        # 窗=D-20..D0 (21 交易日): 第-22日(窗外)不计, 第-21日(=D-20)计入
        out, _, _ = trained
        monkeypatch.setattr(
            fbp,
            "_LHB_DAILY_CACHE",
            self._daily(
                [
                    ("600600", DATES[-22], 1.0, 5.0, 1.0, 1.0),
                    ("600600", DATES[-21], 1.0, 1.0, 1.0, 0.0),
                    ("600600", DATES[-1], 1.0, 1.0, 1.0, 0.0),
                ]
            ),
        )
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})
        page = fbp.serve_firstboard(_today_events_df(1), models_dir=out)
        r = page.iloc[0]
        assert r["LHB上榜日数(20日)"] == 2.0
        assert r["LHB净买亿(20日)"] == pytest.approx(2.0)  # 窗外 5 亿不计

    def test_historical_d0_window(self, monkeypatch):
        monkeypatch.setattr(
            fbp,
            "_LHB_DAILY_CACHE",
            self._daily(
                [
                    ("600600", DATES[-10], 1.0, 1.0, 2.0, 1.0),  # 指定 d0
                    ("600600", DATES[-30], 1.0, 5.0, 5.0, 0.0),  # D-20 → 窗内第1日
                    ("600600", DATES[-31], 1.0, 9.0, 9.0, 0.0),  # D-21 → 窗外
                ]
            ),
        )
        ann = fbp._lhb_annotations(
            pd.DataFrame({"date": DATES}), ["600600"], d0=DATES[-10]
        )
        a = ann["600600"]
        assert a["pre_days"] == 2.0 and a["pre_net"] == 6.0 and a["pre_hot"] == 7.0
        assert a["d0_net"] == 1.0 and a["d0_hot"] == 2.0 and a["pre_inst"] == 1.0

    def test_missing_file_blank_not_fatal(self, trained, monkeypatch):
        # 数据缺失 → 全部标注列空, 行数/既有列不受影响 (fail-soft 契约)
        out, _, _ = trained
        monkeypatch.setattr(fbp, "_LHB_DAILY_CACHE", pd.DataFrame())
        monkeypatch.setattr(fbp, "_NAME_CACHE", {})
        page = fbp.serve_firstboard(_today_events_df(2), models_dir=out)
        assert len(page) == 2
        for c in fbp.LHB_ANNOT_COLS:
            if c == "机构在场(20日)":
                assert (page[c] == "").all()
            else:
                assert page[c].isna().all()


class TestGeniousMultiSheet:
    def test_write_xlsx_extra_sheets(self, tmp_path, monkeypatch):
        genious = importlib.import_module("scripts._genious_excel")
        s1 = pd.DataFrame({"排名": [1], "symbol": ["600001"]})
        extra = pd.DataFrame({"代码": ["600002"], "续板概率5日": [0.42]})
        fp = genious.write_xlsx(
            s1,
            "20990101",
            list_dir=str(tmp_path),
            extra_sheets=[("首板点名", extra, genious.FB_BANNER, genious.FB_LEGEND)],
        )
        import openpyxl

        wb = openpyxl.load_workbook(fp)
        assert wb.sheetnames == ["冠军四段", "首板点名"]
        assert wb["首板点名"]["A1"].value == genious.FB_BANNER

    def test_preboard_sheet_footnote_below_table(self, tmp_path):
        genious = importlib.import_module("scripts._genious_excel")
        pb = pd.DataFrame(
            {
                "通道优先": ["优先", ""],
                "点火旗": ["T1+T2", ""],
                "击中日期": ["2026-09-18", "2026-09-15"],
                "次数10日": [1, 5],
                "次数20日": [2, 5],
                "代码": ["600503", "600501"],
                "当日涨幅": [0.03, np.nan],
                "获利盘": [0.70, 0.42],
                "20日获利盘Δ": [0.12, 0.12],
                "5日均换手": [0.4, 0.4],
                "5日/20日换手比": [0.5, 0.5],
            }
        )
        s1 = pd.DataFrame({"排名": [1], "symbol": ["600001"]})
        fp = genious.write_xlsx(
            s1,
            "20990101",
            list_dir=str(tmp_path),
            extra_sheets=[("板前哨", pb, genious.PB_BANNER, genious.PB_LEGEND)],
        )
        import openpyxl

        wb = openpyxl.load_workbook(fp)
        assert "板前哨" in wb.sheetnames
        ws = wb["板前哨"]
        # 表格数据止于 row 3+len; 脚注=PB_LEGEND 写在其下方 A 列
        foot = [
            ws.cell(row=r, column=1).value
            for r in range(4 + len(pb) + 1, ws.max_row + 1)
        ]
        assert any(isinstance(v, str) and "通道优先" in v for v in foot)
        assert any(isinstance(v, str) and "通道优先·白话" in v for v in foot)
