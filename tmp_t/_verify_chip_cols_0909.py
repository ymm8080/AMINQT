# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

S = Path(r"D:\AMINQT\DAILY OPERATION\STOCK LIST")
wb = load_workbook(
    S / "STOCK LIST 20260909__M20260908__D20260830excessfix_perhorizon.xlsx",
    read_only=True,
)
ws = wb["短名单 T-10"]
rows = list(ws.iter_rows(values_only=True))
hdr = list(rows[0])
i_s, i_wr, i_fl = hdr.index("symbol"), hdr.index("chip_wr5"), hdr.index("chip_flag")
print("xlsx 尾列:", hdr[-3:])
for r in rows[1:]:
    if str(r[i_s]) == "603059":
        print(f"xlsx 603059: chip_wr5={r[i_wr]:.4f} chip_flag={r[i_fl]}")
print(f"xlsx 全表 chip_flag 非空 {sum(1 for r in rows[1:] if r[i_fl])} 行")
wb.close()

doc = Document(str(S / "STOCK LIST 20260909__M20260908__D20260830excessfix.docx"))
for t in doc.tables:
    hdrs = [c.text for c in t.rows[0].cells]
    if "chip_flag" not in hdrs:
        continue
    syms = [t.rows[i].cells[1].text for i in range(1, len(t.rows))]
    if "603059" in syms:
        j = syms.index("603059")
        k_wr, k_fl = hdrs.index("chip_wr5"), hdrs.index("chip_flag")
        print(f"docx 603059: chip_wr5={t.rows[j+1].cells[k_wr].text} "
              f"chip_flag={t.rows[j+1].cells[k_fl].text}")
